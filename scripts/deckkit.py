"""deckkit -- the one shared module for deck generation (2026-09-11).

Ends the copy-paste era: every deck script before this one redefined ~60 lines
of new()/bullets()/pic()/fig_slide() closures, and the TDTClosedLoop branch of
that family inserted figures WIDTH-ONLY at hand-tuned offsets, so tall figures
overflowed the slide (6 of 10 figure slides in build_weekly_deck_2026-09-04.py)
and every one needed hand-editing before presentation.

Two halves, one geometry contract:

  matplotlib half   figures are BORN slide-shaped: new_fig("FULL") yields a
                    canvas whose aspect matches the slide content box, so
                    fit_picture never has to rescue it.
  pptx half         Deck class + fit_picture (aspect-preserving fit with
                    horizontal AND vertical centering -- ported from
                    NNController build_cross_study_synthesis.py and extended).

House style (the professor's rules, encoded once): Arial everywhere, BLACK bold
titles (never blue), 11 pt body, ~1 figure per content slide, numbers read from
data at build time; the build FAILS on a missing artifact (need()/jload()).

Figure color doctrine (dataviz method, validated reference palette):
entity colors are FIXED across every figure in every deck -- MPC is always
blue, Choi always orange, NN always aqua; references/targets are ink grey;
grids are recessive hairlines; no dual axes; top/right spines off.

Import pattern from analysis dirs:
    sys.path.insert(0, str(Path(__file__).resolve().parents[N] / "scripts"))
    import deckkit as dk

Dependencies: python-pptx + matplotlib + stdlib. Deliberately NO PIL --
png_size reads the IHDR bytes directly.
"""
from __future__ import annotations

import json
import struct
import sys
from contextlib import contextmanager
from pathlib import Path

# ---------------------------------------------------------------------------
# Geometry contract (inches). Everything both halves know about the slide.
# ---------------------------------------------------------------------------
SLIDE_W, SLIDE_H = 13.333, 7.5          # 16:9
MARGIN_X = 0.45
TITLE_TOP, TITLE_H = 0.25, 0.60
SUBTITLE_TOP, SUBTITLE_H = 0.85, 0.30
CONTENT_TOP = 1.15                       # no subtitle
CONTENT_TOP_SUB = 1.55                   # with subtitle
CONTENT_BOTTOM = 7.20
CAPTION_H = 0.40
FOOTER_TOP = 7.22
CONTENT_W = 12.4                         # SLIDE_W - 2*MARGIN_X, rounded down

# House style
FONT = "Arial"
TITLE_PT, SUBTITLE_PT, BODY_PT, CAPTION_PT, TABLE_PT = 24, 12, 11, 10, 10
SECTION_PT = 32
FOOTER_PT = 8

# Ink (validated reference palette, light surface)
INK = "#0b0b0b"
INK_2 = "#52514e"
MUTED = "#898781"
GRID = "#e1e0d9"
AXIS = "#c3c2b7"

# Categorical slots, fixed order (validated: adjacent-pair CVD dE >= 8 light).
CAT = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100",
       "#e87ba4", "#008300", "#4a3aa7", "#e34948"]

# Entity colors -- color follows the entity, never the rank. Same in every
# figure of every deck from now on.
COLOR = {
    "mpc": CAT[0],       # blue
    "choi": CAT[1],      # orange
    "nn": CAT[2],        # aqua
    "ref": INK_2,        # reference/target trace: secondary ink, dashed
    "null": MUTED,       # shuffled/sham nulls
    "touch": CAT[6],     # violet -- touch/template identity where needed
    "stim": CAT[7],      # red -- stim/actuator annotations
}

FIGSIZE = {
    # preset: (w_in, h_in) -- aspect matched to the pptx content boxes so the
    # placed picture fills its box instead of being letterboxed.
    "FULL": (12.4, 5.4),     # fig_slide with subtitle + caption
    "WIDE": (12.4, 4.2),     # figure above a bullet block
    "TALL": (8.6, 5.4),      # square-ish content (galleries, matrices)
    "HALF": (6.05, 5.0),     # two_fig_slide panels
}
DPI = 200


# ---------------------------------------------------------------------------
# Small shared utilities
# ---------------------------------------------------------------------------
def need(path) -> Path:
    """Anti-drift gate: the build FAILS loudly on a missing artifact."""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"deck build requires missing artifact: {p}")
    return p


def jload(path) -> dict:
    """Strict JSON load -- no silent fallback (that IS the anti-drift property)."""
    return json.loads(need(path).read_text(encoding="utf-8"))


def png_size(path) -> tuple[int, int] | None:
    """(width, height) px from the PNG IHDR chunk; None if not a PNG."""
    try:
        with open(path, "rb") as f:
            head = f.read(24)
        if head[:8] != b"\x89PNG\r\n\x1a\n":
            return None
        w, h = struct.unpack(">II", head[16:24])
        return int(w), int(h)
    except OSError:
        return None


def save_versioned(prs, out: Path, max_v: int = 9) -> Path:
    """Save; if the deck is open in PowerPoint, write _vN instead (N<=max_v)."""
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(out))
        return out
    except PermissionError:
        n = 2
        while True:
            alt = out.with_name(f"{out.stem}_v{n}{out.suffix}")
            try:
                prs.save(str(alt))
                print(f"NOTE: {out.name} is open in PowerPoint; wrote {alt.name} instead.")
                return alt
            except PermissionError:
                n += 1
                if n > max_v:
                    raise


# ---------------------------------------------------------------------------
# pptx half
# ---------------------------------------------------------------------------
def _pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    return Presentation, Inches, Pt, RGBColor, PP_ALIGN


def _rgb(hexstr):
    from pptx.dml.color import RGBColor
    h = hexstr.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def fit_picture(slide, image, left=MARGIN_X, top=CONTENT_TOP,
                max_w=CONTENT_W, max_h=None, center_h=True, center_v=True):
    """Aspect-preserving fit of a PNG into a declared box, centered both ways.

    The donor logic (NNController fit_picture) plus vertical centering.
    Returns the placed shape so callers/tests can inspect its geometry.
    """
    from pptx.util import Inches
    image = str(need(image))
    if max_h is None:
        max_h = CONTENT_BOTTOM - top
    sz = png_size(image)
    if sz:
        w_px, h_px = sz
        ar = w_px / h_px
        width = max_w
        height = width / ar
        if height > max_h:
            height = max_h
            width = height * ar
    else:  # non-PNG fallback: fill the box (should not happen in our decks)
        width, height = max_w, max_h
    lft = left + (max_w - width) / 2 if center_h else left
    tp = top + (max_h - height) / 2 if center_v else top
    return slide.shapes.add_picture(image, Inches(lft), Inches(tp),
                                    width=Inches(width), height=Inches(height))


class Deck:
    """A 16:9 presentation in the house style. One instance per deck build."""

    def __init__(self, footer: str = ""):
        Presentation, Inches, Pt, RGBColor, PP_ALIGN = _pptx()
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self._blank = self.prs.slide_layouts[6]
        self._footer = footer
        self._n = 0

    # -- low-level text ----------------------------------------------------
    def _text(self, slide, text, left, top, width, height, size, bold=False,
              color=INK, align=None, wrap=True):
        _, Inches, Pt, _, PP_ALIGN = _pptx()
        box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                       Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = wrap
        p = tf.paragraphs[0]
        if align == "center":
            p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = _rgb(color)
        return box

    def _chrome(self, slide, title, subtitle=None):
        self._text(slide, title, MARGIN_X, TITLE_TOP, CONTENT_W, TITLE_H,
                   TITLE_PT, bold=True, color=INK)
        if subtitle:
            self._text(slide, subtitle, MARGIN_X, SUBTITLE_TOP, CONTENT_W,
                       SUBTITLE_H, SUBTITLE_PT, color=INK_2)
        if self._footer:
            self._text(slide, f"{self._footer}  ·  {self._n}", MARGIN_X,
                       FOOTER_TOP, CONTENT_W, 0.25, FOOTER_PT, color=MUTED)

    def _new(self):
        self._n += 1
        return self.prs.slides.add_slide(self._blank)

    # -- slide builders ----------------------------------------------------
    def title_slide(self, title, subtitle=None, lines=()):
        s = self._new()
        self._text(s, title, MARGIN_X, 2.6, CONTENT_W, 1.0, 34, bold=True)
        if subtitle:
            self._text(s, subtitle, MARGIN_X, 3.7, CONTENT_W, 0.5, 16, color=INK_2)
        for i, ln in enumerate(lines):
            self._text(s, ln, MARGIN_X, 4.5 + 0.32 * i, CONTENT_W, 0.3,
                       BODY_PT, color=INK_2)
        return s

    def section_slide(self, title, kicker=None):
        s = self._new()
        if kicker:
            self._text(s, kicker.upper(), MARGIN_X, 2.9, CONTENT_W, 0.4, 13,
                       bold=True, color=MUTED)
        self._text(s, title, MARGIN_X, 3.3, CONTENT_W, 1.0, SECTION_PT, bold=True)
        if self._footer:
            self._text(s, f"{self._footer}  ·  {self._n}", MARGIN_X, FOOTER_TOP,
                       CONTENT_W, 0.25, FOOTER_PT, color=MUTED)
        return s

    def slide(self, title, subtitle=None):
        s = self._new()
        self._chrome(s, title, subtitle)
        return s

    def bullets(self, slide, items, top=None, size=BODY_PT, left=0.72,
                width=12.0, line_h=None):
        """items: iterable of (level, text); level 0 black, level 1 grey."""
        _, Inches, Pt, _, _ = _pptx()
        top = CONTENT_TOP if top is None else top
        line_h = line_h or (size + 7) / 72.0 * 1.35
        box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                       Inches(width),
                                       Inches(min(CONTENT_BOTTOM - top,
                                                  line_h * max(len(items), 1))))
        tf = box.text_frame
        tf.word_wrap = True
        for i, (lvl, text) in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.level = lvl
            p.space_after = Pt(4)
            r = p.add_run()
            r.text = ("• " if lvl == 0 else "– ") + text
            r.font.name = FONT
            r.font.size = Pt(size if lvl == 0 else size - 1)
            r.font.bold = False
            r.font.color.rgb = _rgb(INK if lvl == 0 else INK_2)
        return box

    def bullets_slide(self, title, items, subtitle=None, size=BODY_PT):
        s = self.slide(title, subtitle)
        self.bullets(s, items,
                     top=CONTENT_TOP_SUB if subtitle else CONTENT_TOP, size=size)
        return s

    def fig_slide(self, title, png, subtitle=None, caption=None,
                  bullets_below=None):
        s = self.slide(title, subtitle)
        top = CONTENT_TOP_SUB if subtitle else CONTENT_TOP
        bottom = CONTENT_BOTTOM
        if caption:
            bottom -= CAPTION_H
        if bullets_below:
            bottom -= 0.30 * len(bullets_below) + 0.15
        fit_picture(s, png, top=top, max_h=bottom - top)
        y = bottom
        if bullets_below:
            self.bullets(s, bullets_below, top=y + 0.10)
            y += 0.30 * len(bullets_below) + 0.15
        if caption:
            self._text(s, caption, MARGIN_X, CONTENT_BOTTOM - CAPTION_H + 0.05,
                       CONTENT_W, CAPTION_H, CAPTION_PT, color=INK_2,
                       align="center")
        return s

    def two_fig_slide(self, title, png_l, png_r, subtitle=None,
                      captions=(None, None)):
        s = self.slide(title, subtitle)
        top = CONTENT_TOP_SUB if subtitle else CONTENT_TOP
        gap = 0.30
        half_w = (CONTENT_W - gap) / 2
        bottom = CONTENT_BOTTOM - (CAPTION_H if any(captions) else 0)
        for png, cap, left in ((png_l, captions[0], MARGIN_X),
                               (png_r, captions[1], MARGIN_X + half_w + gap)):
            fit_picture(s, png, left=left, top=top, max_w=half_w,
                        max_h=bottom - top)
            if cap:
                self._text(s, cap, left, CONTENT_BOTTOM - CAPTION_H + 0.05,
                           half_w, CAPTION_H, CAPTION_PT, color=INK_2,
                           align="center")
        return s

    def table_slide(self, title, headers, rows, subtitle=None, col_widths=None,
                    highlight_rows=(), size=TABLE_PT):
        _, Inches, Pt, _, _ = _pptx()
        s = self.slide(title, subtitle)
        top = CONTENT_TOP_SUB if subtitle else CONTENT_TOP
        n_r, n_c = len(rows) + 1, len(headers)
        height = min(CONTENT_BOTTOM - top, 0.32 * n_r)
        shp = s.shapes.add_table(n_r, n_c, Inches(MARGIN_X), Inches(top),
                                 Inches(CONTENT_W), Inches(height))
        tbl = shp.table
        if col_widths:
            total = sum(col_widths)
            for j, wfrac in enumerate(col_widths):
                tbl.columns[j].width = Inches(CONTENT_W * wfrac / total)
        for j, h in enumerate(headers):
            self._cell(tbl.cell(0, j), str(h), size, bold=True)
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                self._cell(tbl.cell(i, j), str(v), size,
                           bold=(i - 1) in highlight_rows)
        return s

    def _cell(self, cell, text, size, bold=False):
        _, _, Pt, _, _ = _pptx()
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = _rgb(INK)

    def kv_band(self, slide, stats, top=None):
        """stats: [(value, label)] -- N boxes across CONTENT_W."""
        top = (CONTENT_BOTTOM - 1.1) if top is None else top
        n = len(stats)
        w = CONTENT_W / n
        for i, (value, label) in enumerate(stats):
            left = MARGIN_X + i * w
            self._text(slide, str(value), left, top, w - 0.1, 0.5, 22,
                       bold=True, align="center")
            self._text(slide, str(label), left, top + 0.55, w - 0.1, 0.35,
                       CAPTION_PT, color=INK_2, align="center")

    def save(self, out) -> Path:
        return save_versioned(self.prs, Path(out))


# ---------------------------------------------------------------------------
# matplotlib half
# ---------------------------------------------------------------------------
@contextmanager
def deck_style():
    """rc context: Arial (with fallback), slide-legible sizes, recessive chrome."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    rc = {
        "font.family": "sans-serif",
        "font.sans-serif": ["Arial", "DejaVu Sans"],
        "figure.dpi": 100,
        "savefig.dpi": DPI,
        "figure.constrained_layout.use": True,
        "axes.titlesize": 13,
        "axes.titleweight": "bold",
        "axes.labelsize": 11,
        "xtick.labelsize": 10,
        "ytick.labelsize": 10,
        "legend.fontsize": 10,
        "axes.edgecolor": AXIS,
        "axes.labelcolor": INK,
        "text.color": INK,
        "xtick.color": INK_2,
        "ytick.color": INK_2,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.grid": True,
        "grid.color": GRID,
        "grid.linewidth": 0.6,
        "grid.alpha": 1.0,
        "axes.axisbelow": True,
        "lines.linewidth": 2.0,
        "axes.prop_cycle": matplotlib.cycler(color=CAT),
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.facecolor": "white",
    }
    with plt.rc_context(rc):
        yield plt


def new_fig(preset="FULL", nrows=1, ncols=1, **kw):
    """Figure born slide-shaped. Use inside deck_style()."""
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(nrows, ncols, figsize=FIGSIZE[preset], **kw)
    return fig, ax


def save_fig(fig, path) -> Path:
    """Save at deck dpi WITHOUT bbox_inches='tight' -- tight would break the
    declared aspect; constrained_layout already handles spacing."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(str(path), dpi=DPI)
    import matplotlib.pyplot as plt
    plt.close(fig)
    sz = png_size(path)
    print(f"wrote {path} ({sz[0]}x{sz[1]} px)" if sz else f"wrote {path}")
    return path


if __name__ == "__main__":
    # Smoke: build a 4-slide demo deck + a demo figure, then print geometry.
    out_dir = Path(__file__).resolve().parent.parent / "outputs" / "deck_figures"
    with deck_style() as plt:
        fig, ax = new_fig("FULL")
        ax.plot([0, 1, 2, 3], [0, 1, 0.5, 1.2], label="MPC", color=COLOR["mpc"])
        ax.plot([0, 1, 2, 3], [0, 0.8, 0.6, 0.9], label="Choi", color=COLOR["choi"])
        ax.set_xlabel("run")
        ax.set_ylabel("r")
        ax.legend(frameon=False)
        demo = save_fig(fig, out_dir / "deckkit_smoke.png")
    d = Deck(footer="deckkit smoke")
    d.title_slide("deckkit smoke test", "self-test artifact", ["generated by deckkit.py __main__"])
    d.section_slide("Section divider", kicker="smoke")
    d.fig_slide("FULL figure fits its box", demo, subtitle="subtitle band",
                caption="caption band")
    d.table_slide("Table", ["a", "b"], [["1", "2"], ["3", "4"]])
    p = d.save(out_dir / "deckkit_smoke.pptx")
    print(f"smoke deck: {p}")
